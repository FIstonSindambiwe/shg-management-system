"""
Excel files: the report as a workbook, and the empty template used for importing.
"""
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import Transaction

T = Transaction.Type
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="E5E7EB")
THIN = Border(bottom=Side(style="thin", color="9CA3AF"))
MONEY = "#,##0"


def _sheet(ws, headers, rows, money_cols=(), widths=None, total=False):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=ws.max_row, column=c)
        cell.font = BOLD
        cell.fill = HEAD_FILL
        cell.border = THIN
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    if total and rows:
        ws.append(["Byose (Total)"] + [
            sum((r[i] or 0) for r in rows) if (i + 1) in money_cols else "" for i in range(1, len(headers))
        ])
        for c in range(1, len(headers) + 1):
            ws.cell(row=ws.max_row, column=c).font = BOLD
    for c in money_cols:
        for row in ws.iter_rows(min_row=2, min_col=c, max_col=c):
            for cell in row:
                cell.number_format = MONEY
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def build_xlsx(summary, title, generated_by):
    wb = Workbook()
    tot, cash, multi = summary["totals"], summary.get("cash"), len(summary["groups"]) > 1

    ws = wb.active
    ws.title = "Incamake (Summary)"
    ws.append([title]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Igihe (Period): {summary['period']['label']}"])
    ws.append([f"Byakozwe na (Prepared by): {generated_by} - {timezone.localtime():%d/%m/%Y %H:%M}"])
    ws.append([])
    _sheet(ws, ["Ikintu (Item)", "Amafaranga (RWF)"], [
        ["Imigabane (Shares bought)", tot[T.SHARE]], ["Ingoboka (Social fund in)", tot[T.INGOBOKA_IN]],
        ["Ingoboka yatanzwe (Social fund paid out)", tot[T.INGOBOKA_OUT]], ["Mituweli (Mutuelle)", tot[T.MUTUELLE]],
        ["Ibihano (Fines paid)", tot[T.FINE_PAYMENT]], ["Inyungu (Interest)", tot[T.INTEREST]],
        ["Inguzanyo zatanzwe (Loans given)", tot[T.LOAN_ISSUE]], ["Ayishyuwe (Loans returned)", tot[T.LOAN_RETURN]],
        ["Andi mafaranga yinjiye (Other money in)", tot[T.OTHER_INCOME]], ["Ibyakoreshejwe (Money spent)", tot[T.EXPENSE]],
        ["Byajyanywe kuri banki (Taken to the bank)", tot[T.BANK_DEPOSIT]],
        ["Inguzanyo zitarishyurwa uyu munsi (Loans still owed today)", summary["loans"]["outstanding"]],
    ], money_cols=(2,), widths=[52, 20])
    ws.freeze_panes = None

    if cash:
        ws = wb.create_sheet("Imari (Cash)")
        P, D = cash["period"], cash["to_date"]
        rows = [[f"+ {lab}", P["in"][k], D["in"][k]] for k, lab in [
            (T.SHARE, "Imigabane (Shares)"), (T.INGOBOKA_IN, "Ingoboka (Social fund)"), (T.FINE_PAYMENT, "Ibihano (Fines)"),
            (T.INTEREST, "Inyungu (Interest)"), (T.LOAN_RETURN, "Ayishyuwe (Loans returned)"),
            (T.OTHER_INCOME, "Andi mafaranga yinjiye (Other money in)")]]
        rows.append(["Byinjiye byose (Total in)", P["total_in"], D["total_in"]])
        rows += [[f"- {lab}", P["out"][k], D["out"][k]] for k, lab in [
            (T.LOAN_ISSUE, "Inguzanyo zatanzwe (Loans given)"), (T.INGOBOKA_OUT, "Ingoboka yatanzwe (Social fund paid out)"),
            (T.EXPENSE, "Ibyakoreshejwe (Money spent)")]]
        rows += [["Byasohotse byose (Total out)", P["total_out"], D["total_out"]],
                 ["Igiteranyo (In minus out)", P["net"], D["net"]],
                 ["- Byajyanywe kuri banki (Taken to the bank)", P["bank_deposits"], D["bank_deposits"]],
                 ["Asigaye mu ntoki (Cash still in hand)", P["in_hand"], D["in_hand"]],
                 ["Mituweli ibitswe ukwayo (Mutuelle, kept apart)", P["mutuelle"], D["mutuelle"]]]
        _sheet(ws, ["", "Muri iki gihe (This period)", "Kugeza ubu (Up to the end date)"], rows,
               money_cols=(2, 3), widths=[52, 24, 26])
        for r in (8, 12, 13, 15):
            for c in (1, 2, 3):
                ws.cell(row=r, column=c).font = BOLD

        moves = [e for e in cash["entries"] if e["type"] in (T.OTHER_INCOME, T.EXPENSE, T.BANK_DEPOSIT)]
        if moves:
            ws = wb.create_sheet("Banki n'ibindi (Bank, other)")
            kinds = {T.OTHER_INCOME: "Yinjiye (In)", T.EXPENSE: "Yakoreshejwe (Spent)", T.BANK_DEPOSIT: "Banki (Bank)"}
            _sheet(ws, ["Itariki (Date)", "Itsinda (Group)", "Ubwoko (Type)", "Ibisobanuro (Details)",
                        "Uwayajyanye (Taken by)", "Amafaranga (RWF)"],
                   [[e["date"], e["shg_name"], kinds[e["type"]], e["description"], e["carried_by"], e["amount"]]
                    for e in moves], money_cols=(6,), widths=[14, 18, 18, 40, 24, 16])

    if multi:
        ws = wb.create_sheet("Amatsinda (Groups)")
        _sheet(ws, ["Itsinda (Group)", "Abanyamuryango", "Imigabane", "Ingoboka", "Mituweli", "Ibihano",
                    "Inyungu", "Inguzanyo", "Ayishyuwe"],
               [[g["name"], g["members"], g["totals"][T.SHARE], g["totals"][T.INGOBOKA_IN], g["totals"][T.MUTUELLE],
                 g["totals"][T.FINE_PAYMENT], g["totals"][T.INTEREST], g["totals"][T.LOAN_ISSUE],
                 g["totals"][T.LOAN_RETURN]] for g in summary["by_group"]],
               money_cols=range(3, 10), widths=[22, 16] + [14] * 7, total=True)

    ws = wb.create_sheet("Abanyamuryango (Members)")
    _sheet(ws, ["Nomero (Code)", "Izina (Name)", "Itsinda (Group)", "Imigabane (count)", "Imigabane (RWF)",
                "Ingoboka", "Mituweli", "Ibihano", "Inyungu", "Inguzanyo (given)", "Ayishyuwe (returned)",
                "Imigabane kugeza ubu (count)"],
           [[m["member_code"], m["full_name"], m["shg_name"], m["shares"], m[T.SHARE], m[T.INGOBOKA_IN],
             m[T.MUTUELLE], m[T.FINE_PAYMENT], m[T.INTEREST], m[T.LOAN_ISSUE], m[T.LOAN_RETURN], m["shares_to_date"]]
            for m in summary["members"]],
           money_cols=range(5, 12), widths=[15, 28, 16, 12, 14, 12, 12, 12, 12, 14, 14, 14], total=True)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- Import template ---------------------------------------------------------------------
MEMBER_HEADERS = ["Izina ry'umuryango (Surname)", "Izina (First name)", "Telefoni (Phone)"]
ENTRY_HEADERS = ["Itariki (Date)", "Izina (Name: SURNAME First)", "Imigabane (Shares count)", "Ingoboka",
                 "Mituweli", "Ibihano (Fines paid)", "Ayishyuwe (Loan returned)", "Inyungu (Interest)",
                 "Inguzanyo nshya (New loan)"]


def build_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Members"
    _sheet(ws, MEMBER_HEADERS, [["KABEHO", "Christine", "0788123456"], ["UWASE", "Odette", ""]], widths=[30, 24, 18])
    ws = wb.create_sheet("Entries")
    _sheet(ws, ENTRY_HEADERS, [
        ["05/03/2026", "KABEHO Christine", 10, 100, 500, 0, 0, 0, 0],
        ["05/03/2026", "UWASE Odette", 4, 100, 500, 100, 0, 0, 20000],
        ["12/03/2026", "UWASE Odette", 2, 100, 500, 0, 5000, 1400, 0],
    ], widths=[14, 28, 16, 12, 12, 16, 20, 16, 20])
    ws = wb.create_sheet("Soma (Read me)")
    for line in [
        "Uko wuzuza (How to fill this in)",
        "",
        "Members: one row per member. Surname first, in CAPITALS if you like. Phone is optional.",
        "Entries: one row per member per meeting date. Leave a cell empty or 0 when nothing was paid.",
        "Itariki (Date): dd/mm/yyyy, e.g. 05/03/2026. Every date becomes a meeting.",
        "Izina (Name): exactly as in the Members sheet: SURNAME First, e.g. KABEHO Christine.",
        "Imigabane: the NUMBER of shares (the system multiplies by the share price, e.g. 500).",
        "Ayishyuwe / Inyungu: paid on the member's oldest open loan.",
        "Inguzanyo nshya: a new loan given that day. It is due 3 months later (group setting).",
        "Nothing is saved if any row has a mistake: you get the list of rows to fix.",
    ]:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 100
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()