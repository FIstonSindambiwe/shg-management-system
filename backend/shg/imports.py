"""
Import a filled-in template (see excel.build_template) into one group.

- "Members" sheet: new members are created; names already in the group are reused.
- "Entries" sheet: one row per member per date. Each date becomes a meeting; shares, ingoboka,
  Mutuelle, fines, loan returns and interest are saved like a weekly sheet; new loans are given
  after that day's contributions.
- Everything is checked first. If any row is wrong, nothing is saved and every problem is listed
  with its sheet and row number. dry_run=True does the whole thing and then undoes it (a preview).
"""
import uuid
from datetime import date, datetime

from django.core.exceptions import ValidationError
from django.db import transaction as db
from openpyxl import load_workbook

from . import services
from .models import AuditLog, Cycle, Loan, Meeting, Member, Transaction

T = Transaction.Type


class ImportFailed(Exception):
    def __init__(self, errors, stats):
        super().__init__("import failed")
        self.errors, self.stats = errors, stats


class _DryRun(Exception):
    pass


def _key(name):
    return " ".join(str(name or "").split()).lower()


def _num(v):
    if v in (None, ""):
        return 0
    try:
        n = float(str(v).replace(",", "").strip())
    except ValueError:
        raise ValueError(f"'{v}' is not a number")
    if n < 0:
        raise ValueError("negative amounts are not allowed")
    return int(round(n))


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"'{v}' is not a date (use dd/mm/yyyy)")


def _rows(ws):
    """Data rows with their Excel row numbers, skipping the header and empty rows."""
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row and any(c not in (None, "") for c in row):
            yield i, list(row) + [None] * 12


def _sheet(wb, *names):
    for n in wb.sheetnames:
        if n.strip().lower() in names:
            return wb[n]
    return None

def import_workbook(file, shg, user, dry_run=False):
    try:
        wb = load_workbook(file, data_only=True, read_only=True)
    except Exception:
        raise ImportFailed([{"sheet": "-", "row": "-", "message": "This is not an .xlsx file the system can read."}], {})

    members_ws = _sheet(wb, "members", "abanyamuryango")
    entries_ws = _sheet(wb, "entries", "ibyinjijwe")
    if not members_ws and not entries_ws:
        raise ImportFailed([{"sheet": "-", "row": "-",
                             "message": "No 'Members' or 'Entries' sheet found. Download the template and use it."}], {})

    errors = []
    stats = {"members_created": 0, "members_matched": 0, "meetings": 0, "entries": 0, "loans": 0, "rows": 0}

    def err(sheet, row, message):
        errors.append({"sheet": sheet, "row": row, "message": message})

    try:
        with db.atomic():
            by_name = {_key(m.full_name): m for m in Member.objects.filter(shg=shg)}
            by_name.update({_key(f"{m.first_name} {m.last_name}"): m for m in Member.objects.filter(shg=shg)})

            # ---- members -------------------------------------------------------------------
            new_members = []
            if members_ws:
                for r, row in _rows(members_ws):
                    last, first, phone = (str(row[0] or "").strip(), str(row[1] or "").strip(),
                                          str(row[2] or "").strip())
                    if not last or not first:
                        err("Members", r, "Surname and first name are both needed.")
                        continue
                    k = _key(f"{last.upper()} {first}")
                    if k in by_name:
                        stats["members_matched"] += 1
                        continue
                    new_members.append((r, last, first, phone))
            if new_members:
                prefix = f"{shg.code}-M"
                used = [int(c[len(prefix):]) for c in Member.objects.filter(member_code__startswith=prefix)
                        .values_list("member_code", flat=True) if c[len(prefix):].isdigit()]
                n = max(used) if used else 0
                for r, last, first, phone in new_members:
                    n += 1
                    m = Member(shg=shg, member_code=f"{prefix}{n:03d}", first_name=first, last_name=last, phone=phone)
                    try:
                        m.full_clean()
                        m.save()
                    except ValidationError as e:
                        err("Members", r, "; ".join(services.error_list(e)))
                        continue
                    by_name[_key(m.full_name)] = m
                    by_name[_key(f"{first} {last}")] = m
                    stats["members_created"] += 1

            # ---- entries, grouped by date ------------------------------------------------------
            days = {}
            if entries_ws:
                for r, row in _rows(entries_ws):
                    stats["rows"] += 1
                    try:
                        d = _date(row[0])
                    except ValueError as e:
                        err("Entries", r, str(e)); continue
                    member = by_name.get(_key(row[1]))
                    if not member:
                        err("Entries", r, f"'{row[1]}' is not in the Members sheet or the group.")
                        continue
                    try:
                        vals = {k: _num(row[i]) for i, k in enumerate(
                            ["shares", "ingoboka", "mutuelle", "fine", "returned", "interest", "new_loan"], start=2)}
                    except ValueError as e:
                        err("Entries", r, str(e)); continue
                    days.setdefault(d, []).append((r, member, vals))

            for d in sorted(days):
                try:
                    cycle = Cycle.for_date(shg, d)
                except ValidationError as e:
                    for r, _, _ in days[d]:
                        err("Entries", r, "; ".join(services.error_list(e)))
                    continue
                meeting = Meeting.objects.filter(shg=shg, date=d).first()
                if not meeting:
                    meeting = Meeting(shg=shg, cycle=cycle, date=d, recorded_by=user, notes="Imported from Excel")
                    try:
                        meeting.full_clean()
                        meeting.save()
                        stats["meetings"] += 1
                    except ValidationError as e:
                        for r, _, _ in days[d]:
                            err("Entries", r, "; ".join(services.error_list(e)))
                        continue

                rows, row_numbers, seen = [], [], set()
                for r, member, v in days[d]:
                    if member.id in seen:
                        err("Entries", r, f"{member.full_name} appears twice on {d:%d/%m/%Y}.")
                        continue
                    seen.add(member.id)
                    payments = []
                    if v["returned"] or v["interest"]:
                        loan = (Loan.objects.filter(member=member).open().order_by("issue_date", "id").first())
                        if not loan:
                            err("Entries", r, f"{member.full_name} has no open loan to return money or interest on.")
                            continue
                        payments = [{"loan": loan, "returned": v["returned"], "interest": v["interest"]}]
                    new_fines = ([{"reason": "OTHER", "note": "Byinjijwe muri Excel (imported)", "amount": v["fine"],
                                   "paid_now": True}] if v["fine"] else [])
                    rows.append({"member": member, "present": None, "shares": v["shares"], "ingoboka": v["ingoboka"],
                                 "mutuelle": v["mutuelle"], "loan_payments": payments, "fine_payments": [],
                                 "new_fines": new_fines})
                    row_numbers.append(r)

                if rows:
                    try:
                        sub, _ = services.save_meeting_sheet(meeting=meeting, user=user,
                                                             submission_id=uuid.uuid4(), rows=rows)
                        stats["entries"] += sub.transaction_count
                    except services.SheetErrors as e:
                        for idx, fields in e.errors.items():
                            msgs = [m for ms in fields.values() for m in ms]
                            err("Entries", row_numbers[int(idx)], "; ".join(msgs))
                        continue
                    except ValidationError as e:
                        for r in row_numbers:
                            err("Entries", r, "; ".join(services.error_list(e)))
                        continue

                # New loans after the day's contributions. History is imported as it happened,
                # so the share-limit and cash checks are skipped (and recorded as such).
                for r, member, v in days[d]:
                    if not v["new_loan"]:
                        continue
                    loan = Loan(shg=shg, cycle=cycle, member=member, principal=v["new_loan"], issue_date=d,
                                meeting=meeting, recorded_by=user, purpose="Imported from Excel")
                    loan.skip_limits = True
                    try:
                        loan.full_clean()
                        loan.save()
                        tx = Transaction(shg=shg, cycle=cycle, meeting=meeting, member=member, date=d,
                                         type=T.LOAN_ISSUE, amount=v["new_loan"], loan=loan, recorded_by=user,
                                         description="Imported from Excel")
                        tx.full_clean()
                        tx.save()
                        stats["loans"] += 1
                    except ValidationError as e:
                        err("Entries", r, "New loan: " + "; ".join(services.error_list(e)))

            if errors:
                raise ImportFailed(sorted(errors, key=lambda e: (e["sheet"], e["row"] if isinstance(e["row"], int) else 0)),
                                   stats)
            if dry_run:
                raise _DryRun()
            AuditLog.record(user, AuditLog.Action.IMPORT, shg, {"excel_import": stats})
    except _DryRun:
        return {"dry_run": True, **stats}
    return {"dry_run": False, **stats}